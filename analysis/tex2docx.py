import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
tex = open(src, encoding='utf-8').read()

def clean(s):
    s = s.strip()
    s = re.sub(r'\\textbf\{(.*?)\}', r'\1', s)
    s = re.sub(r'\\mathbf\{(.*?)\}', r'\1', s)
    rep = [(r'\hat\theta', 'θ̂'), (r'\pm', '±'), (r'^\circ', '°'), (r'^{*}', '*'),
           (r'\sigma_t', 'σt'), (r'\lambda', 'λ'), (r'\Delta', 'Δ'), (r'\ ', ' '),
           ('~', ' '), ('$-$', '−'), ('$', ''), ('{', ''), ('}', '')]
    for a, b in rep: s = s.replace(a, b)
    s = re.sub(r'(?<![\w.])-(?=\d)', '−', s)      # signo menos tipográfico
    return re.sub(r'\s+', ' ', s).strip()

def split_cells(row):
    # split on & only outside braces
    out, depth, cur = [], 0, ''
    for ch in row:
        if ch == '{': depth += 1
        if ch == '}': depth -= 1
        if ch == '&' and depth == 0: out.append(cur); cur = ''
        else: cur += ch
    out.append(cur); return out

def parse_row(row):
    cells = []
    for c in split_cells(row):
        c = c.strip()
        m = re.match(r'\\multicolumn\{(\d+)\}\{[^}]*\}\{(.*)\}$', c, re.S)
        if m: cells.append((clean(m.group(2)), int(m.group(1)), 'mc')); continue
        m = re.match(r'\\multirow\{(\d+)\}\{[^}]*\}\{(.*)\}$', c, re.S)
        if m: cells.append((clean(m.group(2)), 1, ('mr', int(m.group(1))))); continue
        cells.append((clean(c), 1, None))
    return cells

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol); tcPr.append(sh)

def borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for edge, (sz, val) in kw.items():
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), val); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), '000000'); b.append(e)
    tcPr.append(b)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.5))
for m in ('top_margin', 'bottom_margin'): setattr(sec, m, Cm(1.8))
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10)

title = re.search(r'\\title\{(.*?)\}', tex).group(1)
author = re.search(r'\\author\{(.*?)\}', tex).group(1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(title); r.bold = True; r.font.size = Pt(16)
p = doc.add_paragraph(author); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

body = tex.split(r'\newpage', 1)[1]
tokens = re.split(r'(\\section\{.*?\}|\\begin\{tabular\}.*?\\end\{tabular\}|\\caption\{.*?\}\s*\n\s*\\label|\\paragraph\{\*\}.*?(?=\n\s*\n|\n%))', body, flags=re.S)
ntab = 0
for tk in tokens:
    if not tk or not tk.strip(): continue
    if tk.startswith(r'\section{'):
        doc.add_heading(clean(tk[9:-1]), level=1)
    elif tk.startswith(r'\begin{tabular}'):
        inner = re.sub(r'^\\begin\{tabular\}\{[^}]*\}', '', tk)
        inner = inner.replace(r'\end{tabular}', '')
        lines = [l for l in re.split(r'\\\\', inner)]
        rows, header_rows, mid = [], 0, 0
        for l in lines:
            l2 = l
            if r'\midrule' in l2:
                mid += 1
                if mid == 1: header_rows = len(rows)
                l2 = l2.replace(r'\midrule', '')
            l2 = l2.replace(r'\toprule', '').replace(r'\bottomrule', '')
            if not l2.strip(): continue
            rows.append((parse_row(l2), mid))
        ncol = max(sum(c[1] for c in r) for r, _ in rows)
        t = doc.add_table(rows=len(rows), cols=ncol); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Table Grid'
        mr_pending = {}
        for ri, (cells, midx) in enumerate(rows):
            ci = 0
            for text, span, kind in cells:
                cell = t.cell(ri, ci)
                if span > 1: cell = cell.merge(t.cell(ri, ci + span - 1))
                if isinstance(kind, tuple) and kind[0] == 'mr':
                    last = min(ri + kind[1] - 1, len(rows) - 1)
                    cell = cell.merge(t.cell(last, ci)); mr_pending[ci] = last
                if text or not (ci in mr_pending and ri <= mr_pending[ci]):
                    cell.text = ''
                    run = cell.paragraphs[0].add_run(text)
                    run.font.size = Pt(8)
                    if ri < header_rows: run.bold = True
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri < header_rows or ci > 0) else WD_ALIGN_PARAGRAPH.LEFT
                if ri < header_rows: shade(cell, 'E7E6E6')
                if midx >= 2 and ri >= len(rows) - 1 and 'MAE' in rows[ri][0][0][0]: shade(cell, 'F2F2F2')
                ci += span
        for row in t.rows:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(0); par.paragraph_format.space_before = Pt(0)
        ntab += 1
    elif tk.startswith(r'\caption{'):
        cap = clean(re.sub(r'\\caption\{(.*)\}\s*\n\s*\\label$', r'\1', tk, flags=re.S))
        p = doc.add_paragraph(); r = p.add_run(cap); r.italic = True; r.font.size = Pt(9)
    elif tk.startswith(r'\paragraph{*}'):
        p = doc.add_paragraph(); r = p.add_run('* ' + clean(tk[len(r'\paragraph{*}'):])); r.font.size = Pt(9)
doc.save(dst); print('tablas escritas:', ntab)
